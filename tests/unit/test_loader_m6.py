"""M6 loader 扩展单测：Word / Excel / PDF 表格检测（离线，现场生成测试文件）。"""

from pathlib import Path

import docx
import openpyxl
import pymupdf

from app.ingestion.loader import _table_to_markdown, load_file


def _make_docx(path: Path) -> None:
    d = docx.Document()
    d.add_heading("考勤制度", level=1)
    d.add_paragraph("第一条 迟到扣款：每次 20 元。")
    d.add_paragraph("第二条 旷工扣款：当日工资的 200%。")
    d.save(path)


def _make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "职级薪酬"
    ws.append(["职级", "月薪区间（元）"])
    ws.append(["P3", "13000—20000"])
    ws.append(["P4", "21000—30000"])
    wb.save(path)


def _make_pdf_with_table(path: Path) -> None:
    """pymupdf 画一个带线条的表格 + 正文，pdfplumber 应能检测到表格。

    注意：表格检测与文字语言无关，单元格用英文字符——
    嵌入 TTC 中文字体会干扰 pdfplumber 的线条解析，且 CI（ubuntu）没有该字体。
    """
    doc = pymupdf.open()
    page = doc.new_page()
    # 注意：pymupdf 默认 helv 字体不含中文字形，插入后提取不出来，
    # 测试正文统一用英文（中文提取路径已由 test_smoke 的 msyh 字体测试覆盖）
    page.insert_text((72, 72), "Absenteeism deduction: 200% of daily wage.")
    # 画表格线（pdfplumber 依据线条检测表格）
    x, y = 72, 150
    page.draw_rect(pymupdf.Rect(x, y, x + 180, y + 40))
    page.draw_line(pymupdf.Point(x + 90, y), pymupdf.Point(x + 90, y + 40))
    page.draw_line(pymupdf.Point(x, y + 20), pymupdf.Point(x + 180, y + 20))
    page.insert_text((x + 5, y + 15), "Grade")
    page.insert_text((x + 95, y + 15), "Salary")
    page.insert_text((x + 5, y + 35), "P4")
    page.insert_text((x + 95, y + 35), "21000")
    doc.save(path)
    doc.close()


def test_table_to_markdown():
    md = _table_to_markdown([["a", "b"], ["1", "2"]])
    assert "| a | b |" in md
    assert "| 1 | 2 |" in md
    assert "---" in md  # 表头分隔行


def test_load_docx(tmp_path: Path):
    path = tmp_path / "制度.docx"
    _make_docx(path)
    docs = load_file(path)
    assert len(docs) == 1
    assert "旷工扣款" in docs[0].page_content
    assert docs[0].metadata["file_name"] == "制度.docx"


def test_load_xlsx_marks_table(tmp_path: Path):
    path = tmp_path / "薪酬表.xlsx"
    _make_xlsx(path)
    docs = load_file(path)
    assert len(docs) == 1
    assert docs[0].metadata["chunk_kind"] == "table"
    assert docs[0].metadata["sheet_name"] == "职级薪酬"
    assert "P4" in docs[0].page_content
    assert "21000—30000" in docs[0].page_content


def test_load_pdf_detects_table(tmp_path: Path):
    path = tmp_path / "带表格.pdf"
    _make_pdf_with_table(path)
    docs = load_file(path)
    # 正文至少 1 个 + 表格 chunk 至少 1 个
    assert any("Absenteeism" in d.page_content for d in docs)
    table_docs = [d for d in docs if d.metadata.get("chunk_kind") == "table"]
    assert table_docs, "应检测到表格并输出 table chunk"
    assert any("21000" in d.page_content for d in table_docs)
